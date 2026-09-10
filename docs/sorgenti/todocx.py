"""Converte i due manuali in .docx.

Non e' una conversione generica di HTML: riconosce le strutture che i manuali
usano davvero (passi numerati, binari Android/iPhone, riquadri, elenchi
etichetta/descrizione, matrice dei permessi) e le rende con l'equivalente di
Word. Le etichette letterali dell'app restano in monospaziato con fondino,
come nella pagina.
"""
import pathlib
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCRATCH = pathlib.Path(__file__).parent
DOCS = SCRATCH.parent

SERIF = "Cambria"
SANS = "Segoe UI"
MONO = "Consolas"

VERDE = RGBColor(0x00, 0x6B, 0x45)
GRIGIO = RGBColor(0x4C, 0x4B, 0x45)
NERO = RGBColor(0x14, 0x18, 0x1A)
AMBRA = RGBColor(0x6B, 0x42, 0x09)
ROSSO = RGBColor(0x8A, 0x1F, 0x1F)
BLU = RGBColor(0x17, 0x40, 0x8F)

FONDO_CHIP = "F2F1EC"
FONDO_NOTA = "EAF9F1"
FONDO_WARN = "FBF1DE"
FONDO_STOP = "FBE7E7"
FONDO_TESTA = "EFEDE6"

VOID = {"br", "img", "meta", "link", "input", "hr"}
SALTA = {"script", "style", "title"}


# --------------------------------------------------------------- parsing
class Nodo:
    __slots__ = ("tag", "attrs", "children", "text")

    def __init__(self, tag=None, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.children = []
        self.text = ""

    @property
    def classi(self):
        return self.attrs.get("class", "").split()

    def figli(self, *tag):
        return [c for c in self.children if c.tag in tag]

    def primo(self, tag=None, classe=None):
        for c in self.children:
            if tag and c.tag != tag:
                continue
            if classe and classe not in c.classi:
                continue
            return c
        return None


class Albero(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Nodo("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        n = Nodo(tag, attrs)
        self.stack[-1].children.append(n)
        if tag not in VOID:
            self.stack.append(n)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Nodo(tag, attrs))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        n = Nodo()
        n.text = data
        self.stack[-1].children.append(n)


def leggi(path):
    a = Albero()
    a.feed(path.read_text(encoding="utf-8"))
    return a.root


# --------------------------------------------------------------- testo inline
def raccogli(nodo, stile="", out=None):
    out = [] if out is None else out
    for c in nodo.children:
        if c.tag is None:
            out.append([c.text, stile])
        elif c.tag in SALTA:
            continue
        elif c.tag == "br":
            out.append(["\n", stile])
        else:
            s = stile
            if c.tag in ("strong", "b"):
                s = "b"
            elif c.tag in ("em", "i") and stile != "b":
                s = "i"
            if "tap" in c.classi or c.tag == "code":
                s = "mono"
            elif "step-flag" in c.classi:
                s = "flag"
            raccogli(c, s, out)
    return out


def normalizza(pezzi):
    """Collassa gli spazi come farebbe il browser, senza perdere gli stili."""
    fuori = []
    for testo, stile in pezzi:
        t = re.sub(r"\s+", " ", testo)
        if not t:
            continue
        if (not fuori and t == " ") or (fuori and fuori[-1][0].endswith(" ") and t.startswith(" ")):
            t = t.lstrip()
        if not t:
            continue
        fuori.append([t, stile])
    while fuori and not fuori[0][0].strip():
        fuori.pop(0)
    if fuori:
        fuori[0][0] = fuori[0][0].lstrip()
        fuori[-1][0] = fuori[-1][0].rstrip()
    return [p for p in fuori if p[0]]


def ombreggia_run(run, colore):
    rpr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), colore)
    rpr.append(shd)


def scrivi(par, pezzi, colore=None, dimensione=None):
    for testo, stile in normalizza(pezzi):
        if stile == "flag":
            # Nella pagina e' un bollino in maiuscolo: qui diventa un inciso
            run = par.add_run(" [" + testo.upper() + "]")
            run.font.name = SANS
            run.font.size = Pt(8.5)
            run.bold = True
            run.font.color.rgb = AMBRA
            continue
        run = par.add_run(testo)
        run.font.name = MONO if stile == "mono" else SERIF
        if dimensione:
            run.font.size = Pt(dimensione - (0.5 if stile == "mono" else 0))
        elif stile == "mono":
            run.font.size = Pt(9.5)
        if stile == "b":
            run.bold = True
        if stile == "i":
            run.italic = True
        if stile == "mono":
            ombreggia_run(run, FONDO_CHIP)
        if colore:
            run.font.color.rgb = colore
    return par


def testo_di(nodo):
    return "".join(t for t, _ in normalizza(raccogli(nodo)))


# --------------------------------------------------------------- blocchi Word
class Costruttore:
    def __init__(self, doc):
        self.doc = doc

    def par(self, spazio_dopo=6, spazio_prima=0, indent=None, sospeso=None):
        p = self.doc.add_paragraph()
        f = p.paragraph_format
        f.space_after = Pt(spazio_dopo)
        f.space_before = Pt(spazio_prima)
        if indent is not None:
            f.left_indent = Inches(indent)
        if sospeso is not None:
            f.first_line_indent = Inches(-sospeso)
        return p

    def etichetta(self, testo, colore=GRIGIO, spazio_prima=0, dimensione=8):
        p = self.par(spazio_dopo=2, spazio_prima=spazio_prima)
        run = p.add_run(testo.upper())
        run.font.name = SANS
        run.font.size = Pt(dimensione)
        run.bold = True
        run.font.color.rgb = colore
        return p

    def titolo(self, testo, livello, colore=NERO):
        p = self.doc.add_heading(level=livello)
        p.paragraph_format.space_before = Pt(14 if livello == 1 else 10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(testo)
        run.font.name = SANS
        run.font.color.rgb = colore
        run.bold = True
        run.font.size = Pt({0: 24, 1: 15, 2: 12}[livello])
        return p

    def testo(self, nodo, colore=None, spazio_dopo=8):
        p = self.par(spazio_dopo=spazio_dopo)
        scrivi(p, raccogli(nodo), colore=colore)
        return p

    def passi(self, ol):
        for i, li in enumerate(ol.figli("li"), start=1):
            p = self.par(spazio_dopo=5, indent=0.32, sospeso=0.32)
            num = p.add_run(f"{i}.\t")
            num.bold = True
            num.font.name = SANS
            num.font.color.rgb = VERDE
            scrivi(p, raccogli(li))

    def punti(self, ul):
        for li in ul.figli("li"):
            p = self.par(spazio_dopo=4, indent=0.32, sospeso=0.2)
            p.add_run("\u2022\t").font.name = SANS
            scrivi(p, raccogli(li))

    def riquadro(self, nodo):
        classi = nodo.classi
        fondo = FONDO_NOTA
        colore = VERDE
        if "warn" in classi:
            fondo, colore = FONDO_WARN, AMBRA
        elif "stop" in classi:
            fondo, colore = FONDO_STOP, ROSSO

        tab = self.doc.add_table(rows=1, cols=1)
        tab.alignment = WD_TABLE_ALIGNMENT.LEFT
        cella = tab.cell(0, 0)
        ombreggia_cella(cella, fondo)
        cella.paragraphs[0]._p.getparent().remove(cella.paragraphs[0]._p)

        h3 = nodo.primo("h3")
        if h3 is not None:
            p = cella.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(testo_di(h3))
            run.bold = True
            run.font.name = SANS
            run.font.size = Pt(10.5)
            run.font.color.rgb = colore
        for par in nodo.figli("p"):
            p = cella.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            scrivi(p, raccogli(par), dimensione=10)
        self.par(spazio_dopo=8)

    def elenco_definizioni(self, nodo):
        righe = [d for d in nodo.figli("div")]
        if not righe:
            return
        tab = self.doc.add_table(rows=len(righe), cols=2)
        tab.style = "Table Grid"
        tab.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, riga in enumerate(righe):
            nome = riga.primo("span", "name") or riga.primo("dt")
            desc = riga.primo("span", "desc")
            c0, c1 = tab.cell(i, 0), tab.cell(i, 1)
            c0.width = Inches(1.5)
            c1.width = Inches(4.3)
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(2)
            r = p0.add_run(testo_di(nome) if nome is not None else "")
            r.bold = True
            r.font.name = SANS
            r.font.size = Pt(10)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(2)
            if desc is not None:
                scrivi(p1, raccogli(desc), dimensione=10)
        self.par(spazio_dopo=8)

    def domande(self, nodo):
        for blocco in nodo.figli("div"):
            q = blocco.primo("p", "q")
            if q is not None:
                p = self.par(spazio_dopo=2, spazio_prima=6)
                run = p.add_run(testo_di(q))
                run.bold = True
                run.font.name = SANS
                run.font.size = Pt(10.5)
            for par in blocco.figli("p"):
                if "q" in par.classi:
                    continue
                self.testo(par, colore=GRIGIO, spazio_dopo=4)

    def binari(self, nodo):
        for track in nodo.figli("div"):
            testa = track.primo("div", "track-head")
            nome = testa_nome(testa)
            self.etichetta(nome, colore=VERDE, spazio_prima=8, dimensione=9)
            req = track.primo("p", "req")
            if req is not None:
                self.testo(req, colore=GRIGIO, spazio_dopo=4)
            ol = track.primo("ol")
            if ol is not None:
                self.passi(ol)

    def matrice(self, tabella):
        intestazioni = [testo_di(th) for th in tabella.primo("thead").primo("tr").figli("th")]
        corpo = tabella.primo("tbody")
        righe = corpo.figli("tr")
        tab = self.doc.add_table(rows=len(righe) + 1, cols=len(intestazioni))
        tab.style = "Table Grid"
        for j, testa in enumerate(intestazioni):
            cella = tab.cell(0, j)
            ombreggia_cella(cella, FONDO_TESTA)
            p = cella.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if j:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(testa)
            r.bold = True
            r.font.name = SANS
            r.font.size = Pt(8.5)
        for i, riga in enumerate(righe, start=1):
            celle = riga.figli("th", "td")
            for j, sorgente in enumerate(celle):
                cella = tab.cell(i, j)
                p = cella.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                valore = testo_di(sorgente)
                if j:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(valore)
                r.font.name = SANS if j else SERIF
                r.font.size = Pt(10 if j else 9.5)
                if valore == "\u2713":
                    r.font.color.rgb = VERDE
                    r.bold = True
                elif valore == "\u2014":
                    r.font.color.rgb = RGBColor(0xA0, 0x9E, 0x96)
        self.par(spazio_dopo=8)


def testa_nome(testa):
    if testa is None:
        return ""
    h3 = testa.primo("h3")
    os = testa.primo("span", "os")
    parti = [testo_di(h3) if h3 is not None else ""]
    if os is not None:
        parti.append(f"({testo_di(os)})")
    return " ".join(x for x in parti if x)


def ombreggia_cella(cella, colore):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), colore)
    cella._tc.get_or_add_tcPr().append(shd)


def numeri_di_pagina(sezione):
    p = sezione.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name = SANS
    run.font.size = Pt(8)
    inizio = OxmlElement("w:fldChar")
    inizio.set(qn("w:fldCharType"), "begin")
    istruzione = OxmlElement("w:instrText")
    istruzione.set(qn("xml:space"), "preserve")
    istruzione.text = "PAGE"
    fine = OxmlElement("w:fldChar")
    fine.set(qn("w:fldCharType"), "end")
    run._r.append(inizio)
    run._r.append(istruzione)
    run._r.append(fine)


# --------------------------------------------------------------- documento
def converti(sorgente, destinazione, titolo):
    radice = leggi(sorgente)
    doc = Document()

    normale = doc.styles["Normal"]
    normale.font.name = SERIF
    normale.font.size = Pt(10.5)
    normale.paragraph_format.space_after = Pt(8)
    rpr = normale.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), SERIF)

    sezione = doc.sections[0]
    sezione.top_margin = Inches(0.7)
    sezione.bottom_margin = Inches(0.7)
    sezione.left_margin = Inches(0.8)
    sezione.right_margin = Inches(0.8)
    numeri_di_pagina(sezione)

    b = Costruttore(doc)

    def trova(nodo, tag=None, classe=None):
        for c in nodo.children:
            if c.tag in SALTA:
                continue
            if (tag is None or c.tag == tag) and (classe is None or classe in c.classi):
                return c
            trovato = trova(c, tag, classe)
            if trovato is not None:
                return trovato
        return None

    testata = trova(radice, "header", "masthead")
    kicker = testata.primo("span", "kicker")
    if kicker is not None:
        b.etichetta(testo_di(kicker), colore=VERDE, dimensione=9)
    b.titolo(testo_di(testata.primo("h1")), 0)
    standfirst = testata.primo("p", "standfirst")
    if standfirst is not None:
        p = b.par(spazio_dopo=8)
        scrivi(p, raccogli(standfirst), colore=GRIGIO, dimensione=12)
    indirizzo = testata.primo("span", "addressbar")
    if indirizzo is not None:
        p = b.par(spazio_dopo=10)
        r = p.add_run(testo_di(indirizzo))
        r.font.name = MONO
        r.font.size = Pt(10)
        ombreggia_run(r, FONDO_CHIP)
    ruoli = testata.primo("div", "roles")
    if ruoli is not None:
        p = b.par(spazio_dopo=10)
        r = p.add_run(" \u00b7 ".join(testo_di(x) for x in ruoli.figli("span")))
        r.font.name = SANS
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = VERDE

    principale = trova(radice, "main")
    for sezione_html in principale.children:
        if sezione_html.tag == "section":
            rendi_sezione(b, sezione_html)
        elif sezione_html.tag == "p" and "footer" in sezione_html.classi:
            b.par(spazio_dopo=2)
            b.testo(sezione_html, colore=GRIGIO)

    doc.core_properties.title = titolo
    doc.core_properties.language = "it-IT"
    destinazione.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destinazione)
    return destinazione


def rendi_sezione(b, sezione):
    for nodo in sezione.children:
        if nodo.tag is None or nodo.tag in SALTA:
            continue
        classi = nodo.classi
        if nodo.tag == "h2":
            occhiello = nodo.primo("span", "num") or nodo.primo("span", "eyebrow")
            if occhiello is not None:
                b.etichetta(testo_di(occhiello), colore=VERDE, spazio_prima=12, dimensione=9)
            resto = "".join(
                t for t, _ in normalizza([
                    x for x in raccogli(nodo)
                    if occhiello is None or t_non_in(x, testo_di(occhiello))
                ])
            )
            b.titolo(resto.strip(), 1)
        elif nodo.tag == "h3":
            b.titolo(testo_di(nodo), 2)
        elif nodo.tag == "p":
            b.testo(nodo, colore=GRIGIO if "lede" in classi else None)
        elif nodo.tag == "ol" and "steps" in classi:
            b.passi(nodo)
        elif nodo.tag == "ul":
            b.punti(nodo)
        elif nodo.tag == "div" and "tracks" in classi:
            b.binari(nodo)
        elif nodo.tag == "div" and "note" in classi:
            b.riquadro(nodo)
        elif nodo.tag == "div" and "legend" in classi:
            b.elenco_definizioni(nodo)
        elif nodo.tag == "div" and ("faq" in classi or "gotchas" in classi):
            b.domande(nodo)
        elif nodo.tag == "div" and "matrix-scroll" in classi:
            tabella = nodo.primo("table")
            if tabella is not None:
                b.matrice(tabella)


def t_non_in(pezzo, occhiello):
    return pezzo[0].strip() not in occhiello


if __name__ == "__main__":
    for slug, titolo in (
        ("manuale-giocatore", "InCampo - Manuale giocatori"),
        ("manuale-staff", "InCampo - Manuale staff"),
    ):
        out = converti(SCRATCH / f"{slug}.html", DOCS / f"{titolo}.docx", titolo)
        print(out.name, "->", f"{out.stat().st_size // 1024} KB")
